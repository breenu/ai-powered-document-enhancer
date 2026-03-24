"""Document export utilities: DOCX and PDF generation.

Provides DOCX export via python-docx (leveraging DocumentFormatter templates)
and PDF export via fpdf2, with metadata embedding for both formats.
"""

import logging
import os
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Dict, Optional

from docx import Document as DocxDocument
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from fpdf import FPDF

from app.core.document_formatter import DocumentFormatter, FormatTemplate, DEFAULT_TEMPLATES
from app.models.document import Document

logger = logging.getLogger(__name__)


@dataclass
class ExportMetadata:
    """Metadata to embed in exported documents."""
    title: str = ""
    author: str = "AI Document Enhancement System"
    subject: str = ""
    keywords: str = ""
    created: Optional[datetime] = None
    doc_type: str = ""
    ocr_confidence: float = 0.0
    readability_score: float = 0.0
    word_count: int = 0

    @classmethod
    def from_document(cls, doc: Document) -> "ExportMetadata":
        return cls(
            title=doc.filename,
            subject=f"Enhanced document - {doc.filename}",
            doc_type=doc.doc_type.value if doc.doc_type else "",
            ocr_confidence=doc.ocr_confidence,
            readability_score=doc.readability_score,
            word_count=doc.get_word_count(),
            created=doc.created_at,
        )


class _EnhancedPDF(FPDF):
    """FPDF subclass with header/footer support."""

    def __init__(self, title: str = "", *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._doc_title = title

    def header(self):
        if self._doc_title:
            self.set_font("Helvetica", "I", 8)
            self.cell(0, 8, self._doc_title, align="L")
            self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, f"Page {self.page_no()}/{{nb}}", align="C")


class DocumentExporter:
    """Exports Document objects to DOCX and PDF files."""

    def __init__(self, formatter: Optional[DocumentFormatter] = None):
        self.formatter = formatter or DocumentFormatter()

    @staticmethod
    def _ensure_dir(path: str) -> None:
        directory = os.path.dirname(path)
        if directory:
            os.makedirs(directory, exist_ok=True)

    def _embed_docx_metadata(self, docx_doc: DocxDocument,
                             metadata: ExportMetadata) -> None:
        """Write metadata into the DOCX core properties."""
        props = docx_doc.core_properties
        if metadata.title:
            props.title = metadata.title
        if metadata.author:
            props.author = metadata.author
        if metadata.subject:
            props.subject = metadata.subject
        if metadata.keywords:
            props.keywords = metadata.keywords
        if metadata.created:
            props.created = metadata.created
        props.comments = (
            f"Document type: {metadata.doc_type} | "
            f"OCR confidence: {metadata.ocr_confidence:.1f}% | "
            f"Readability: {metadata.readability_score:.1f} | "
            f"Words: {metadata.word_count}"
        )

    def export_docx(self, document: Document, output_path: str,
                    doc_type: Optional[str] = None,
                    metadata: Optional[ExportMetadata] = None) -> str:
        """Export a Document to DOCX with formatting and metadata.

        Args:
            document: The processed Document object.
            output_path: Destination .docx file path.
            doc_type: Override document type for template selection.
            metadata: Custom metadata; auto-generated from document if None.

        Returns:
            The output file path.
        """
        self._ensure_dir(output_path)

        text = document.enhanced_text or document.raw_text
        if not text:
            raise ValueError("Document has no text content to export")

        effective_type = doc_type or (
            document.doc_type.value if document.doc_type else None
        )

        docx_doc = self.formatter.apply_template(text, doc_type=effective_type)

        if document.summary_text:
            docx_doc.add_page_break()
            docx_doc.add_heading("Summary", level=1)
            docx_doc.add_paragraph(document.summary_text)

        meta = metadata or ExportMetadata.from_document(document)
        self._embed_docx_metadata(docx_doc, meta)

        docx_doc.save(output_path)
        logger.info("DOCX exported to %s", output_path)
        return output_path

    def export_pdf(self, document: Document, output_path: str,
                   metadata: Optional[ExportMetadata] = None,
                   font_family: str = "Helvetica",
                   font_size: int = 12,
                   line_height: float = 6.0) -> str:
        """Export a Document to PDF via fpdf2.

        Args:
            document: The processed Document object.
            output_path: Destination .pdf file path.
            metadata: Custom metadata; auto-generated from document if None.
            font_family: PDF font family (Helvetica, Courier, Times).
            font_size: Body font size in points.
            line_height: Line height in mm.

        Returns:
            The output file path.
        """
        self._ensure_dir(output_path)

        text = document.enhanced_text or document.raw_text
        if not text:
            raise ValueError("Document has no text content to export")

        meta = metadata or ExportMetadata.from_document(document)

        pdf = _EnhancedPDF(title=meta.title)
        pdf.alias_nb_pages()
        pdf.set_auto_page_break(auto=True, margin=20)

        pdf.set_title(meta.title)
        pdf.set_author(meta.author)
        pdf.set_subject(meta.subject)
        pdf.set_keywords(meta.keywords)
        pdf.set_creator("AI Document Enhancement System")
        pdf.set_creation_date(meta.created or datetime.now())

        pdf.add_page()

        if meta.title:
            pdf.set_font(font_family, "B", font_size + 6)
            pdf.cell(0, 12, meta.title, ln=True, align="C")
            pdf.ln(4)

        pdf.set_font(font_family, "", font_size)

        paragraphs = text.split("\n")
        for para in paragraphs:
            stripped = para.strip()
            if not stripped:
                pdf.ln(line_height)
                continue

            is_heading = (
                stripped.isupper()
                or stripped.endswith(":")
                or (len(stripped.split()) <= 6 and not stripped.endswith("."))
            )

            if is_heading and len(stripped) < 100:
                pdf.ln(line_height / 2)
                pdf.set_font(font_family, "B", font_size + 2)
                pdf.multi_cell(0, line_height, stripped)
                pdf.set_font(font_family, "", font_size)
                pdf.ln(line_height / 4)
            else:
                pdf.multi_cell(0, line_height, stripped)
                pdf.ln(line_height / 3)

        if document.summary_text:
            pdf.add_page()
            pdf.set_font(font_family, "B", font_size + 4)
            pdf.cell(0, 10, "Summary", ln=True)
            pdf.ln(2)
            pdf.set_font(font_family, "", font_size)
            pdf.multi_cell(0, line_height, document.summary_text)

        _add_pdf_stats_footer(pdf, meta, font_family)

        pdf.output(output_path)
        logger.info("PDF exported to %s", output_path)
        return output_path

    def export(self, document: Document, output_path: str,
               doc_type: Optional[str] = None,
               metadata: Optional[ExportMetadata] = None) -> str:
        """Auto-detect format from extension and export accordingly."""
        ext = os.path.splitext(output_path)[1].lower()
        if ext == ".docx":
            return self.export_docx(document, output_path, doc_type, metadata)
        elif ext == ".pdf":
            return self.export_pdf(document, output_path, metadata)
        else:
            raise ValueError(f"Unsupported export format: {ext}. Use .docx or .pdf")


def _add_pdf_stats_footer(pdf: FPDF, meta: ExportMetadata,
                          font_family: str) -> None:
    """Append a small statistics block at the end of the PDF."""
    pdf.ln(8)
    pdf.set_draw_color(180, 180, 180)
    pdf.line(pdf.get_x(), pdf.get_y(), pdf.get_x() + pdf.epw, pdf.get_y())
    pdf.ln(4)
    pdf.set_font(font_family, "I", 8)
    stats = (
        f"Generated by AI Document Enhancement System | "
        f"Type: {meta.doc_type or 'N/A'} | "
        f"OCR Confidence: {meta.ocr_confidence:.1f}% | "
        f"Readability: {meta.readability_score:.1f} | "
        f"Words: {meta.word_count}"
    )
    pdf.multi_cell(0, 4, stats)
